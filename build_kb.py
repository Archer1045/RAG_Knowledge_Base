"""从骑手招聘助手的三个源文件生成 RAG 知识库。"""
from __future__ import annotations

import json
import re
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
DEFAULT_SOURCE_DIR = HERE.parent
DATA_DIR = HERE / "data"
DOCS_DIR = HERE / "documents"
SOURCE_FILES = ("faq.txt", "instructions.txt", "sites_geocoded.xlsx")


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig")


def parse_faq(text: str) -> list[dict[str, str]]:
    entries: list[dict[str, str]] = []
    current: dict[str, str] | None = None
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("问:"):
            if current is not None:
                entries.append(current)
            current = {"question": line[2:].strip(), "answer": ""}
        elif line.startswith("答:") and current is not None:
            current["answer"] = line[2:].strip()
        elif current is not None:
            current["answer"] += " " + line
    if current is not None:
        entries.append(current)
    return entries


def parse_instructions(text: str) -> list[dict[str, str]]:
    items: list[dict[str, str]] = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        match = re.match(r"^([^:：]+)[:：](.*)$", line)
        if match:
            items.append({"key": match.group(1).strip(), "value": match.group(2).strip()})
    return items


def to_float(value):
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def parse_sites(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    if not rows:
        return []
    headers = [str(header).strip() if header is not None else "" for header in rows[0]]
    sites = []
    for values in rows[1:]:
        record = {headers[i]: values[i] for i in range(min(len(headers), len(values)))}
        site = str(record.get("Site") or "").strip()
        district = str(record.get("District") or "").strip()
        if not site or not district:
            continue
        sites.append(
            {
                "city": str(record.get("City") or "").strip(),
                "district": district,
                "site": site,
                "address": str(record.get("Address") or "").strip(),
                "latitude": to_float(record.get("Latitude")),
                "longitude": to_float(record.get("Longitude")),
            }
        )
    return sites


def build_chunks(faq: list[dict], instructions: list[dict], sites: list[dict]) -> list[dict]:
    chunks: list[dict] = []
    for index, item in enumerate(faq, 1):
        chunks.append(
            {
                "id": f"faq-{index:03d}",
                "source": "faq.txt",
                "type": "faq",
                "title": item["question"],
                "content": f"问：{item['question']}\n答：{item['answer']}",
                "metadata": {"question": item["question"], "answer": item["answer"]},
            }
        )
    for index, item in enumerate(instructions, 1):
        chunks.append(
            {
                "id": f"instruction-{index:02d}",
                "source": "instructions.txt",
                "type": "instruction",
                "title": item["key"],
                "content": f"{item['key']}：{item['value']}",
                "metadata": {"topic": item["key"], "value": item["value"]},
            }
        )
    for index, site in enumerate(sites, 1):
        chunks.append(
            {
                "id": f"site-{index:03d}",
                "source": "sites_geocoded.xlsx",
                "type": "site",
                "title": site["site"],
                "content": f"站点：{site['site']}\n区域：{site['district']}\n地址：{site['address']}",
                "metadata": {
                    "city": site["city"],
                    "district": site["district"],
                    "address": site["address"],
                    "latitude": site["latitude"],
                    "longitude": site["longitude"],
                },
            }
        )
    grouped: dict[str, list[str]] = defaultdict(list)
    for site in sites:
        grouped[site["district"]].append(site["site"])
    for district in sorted(grouped):
        site_names = grouped[district]
        chunks.append(
            {
                "id": f"district-{district}",
                "source": "sites_geocoded.xlsx",
                "type": "district",
                "title": f"{district}站点概览",
                "content": f"{district}共有{len(site_names)}个骑手站点，覆盖站点：{'、'.join(site_names)}。",
                "metadata": {"district": district, "site_count": len(site_names), "site_names": site_names},
            }
        )
    return chunks


def faq_markdown(faq: list[dict]) -> str:
    lines = ["# 骑手招聘 FAQ", "", "来源：faq.txt", "", f"共 {len(faq)} 条问答", ""]
    for index, item in enumerate(faq, 1):
        lines.extend(
            [
                f"## {index}. {item['question']}",
                "",
                f"**问：** {item['question']}",
                "",
                f"**答：** {item['answer']}",
                "",
            ]
        )
    return "\n".join(lines).rstrip() + "\n"


def instructions_markdown(instructions: list[dict]) -> str:
    lines = ["# 岗位规则", "", "来源：instructions.txt", "", "| 主题 | 内容 |", "| --- | --- |"]
    for item in instructions:
        value = item["value"].replace("|", "\\|")
        lines.append(f"| {item['key']} | {value} |")
    return "\n".join(lines) + "\n"


def sites_markdown(sites: list[dict]) -> str:
    lines = ["# 北京骑手站点", "", "来源：sites_geocoded.xlsx", "", f"共 {len(sites)} 个站点", ""]
    grouped: dict[str, list[dict]] = defaultdict(list)
    for site in sites:
        grouped[site["district"]].append(site)
    for district in sorted(grouped):
        items = sorted(grouped[district], key=lambda site: site["site"])
        lines.extend([f"## {district}（{len(items)} 个站点）", "", "| 站点 | 地址 | 纬度 | 经度 |", "| --- | --- | --- | --- |"])
        for site in items:
            latitude = f"{site['latitude']:.6f}" if site["latitude"] is not None else ""
            longitude = f"{site['longitude']:.6f}" if site["longitude"] is not None else ""
            lines.append(f"| {site['site']} | {site['address']} | {latitude} | {longitude} |")
        lines.append("")
    return "\n".join(lines) + "\n"


def main() -> None:
    source_dir = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DEFAULT_SOURCE_DIR
    faq = parse_faq(read_text(source_dir / "faq.txt"))
    instructions = parse_instructions(read_text(source_dir / "instructions.txt"))
    sites = parse_sites(source_dir / "sites_geocoded.xlsx")
    chunks = build_chunks(faq, instructions, sites)
    districts = sorted({site["district"] for site in sites})

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    DOCS_DIR.mkdir(parents=True, exist_ok=True)

    index = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "sources": list(SOURCE_FILES),
        "source_dir": str(source_dir),
        "counts": {
            "faq": len(faq),
            "instructions": len(instructions),
            "sites": len(sites),
            "districts": len(districts),
            "chunks": len(chunks),
        },
        "districts": districts,
        "chunk_schema": {
            "id": "string",
            "source": "string",
            "type": "faq|instruction|site|district",
            "title": "string",
            "content": "string",
            "metadata": "object",
        },
    }

    with (DATA_DIR / "chunks.jsonl").open("w", encoding="utf-8") as fh:
        for chunk in chunks:
            fh.write(json.dumps(chunk, ensure_ascii=False) + "\n")
    (DATA_DIR / "documents.json").write_text(json.dumps({"faq": faq, "instructions": instructions, "sites": sites}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (DATA_DIR / "index.json").write_text(json.dumps(index, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (DOCS_DIR / "faq.md").write_text(faq_markdown(faq), encoding="utf-8")
    (DOCS_DIR / "instructions.md").write_text(instructions_markdown(instructions), encoding="utf-8")
    (DOCS_DIR / "sites.md").write_text(sites_markdown(sites), encoding="utf-8")

    print(json.dumps(index["counts"], ensure_ascii=False, indent=2))
    print(f"知识库已生成到: {HERE}")


if __name__ == "__main__":
    main()
